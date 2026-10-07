from pathlib import Path
from typing import Iterable, Union
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument
from app.core.config import get_settings

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Union[Path, str]) -> list[Document]:
    path = Path(str(path))
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()
    if suffix in {".txt", ".md"}:
        return TextLoader(str(path), encoding="utf-8").load()
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")


def load_sample_kb(kb_dir: Union[Path, str, None] = None) -> list[Document]:
    if kb_dir is None:
        kb_dir = Path(get_settings().sample_kb_dir)
    else:
        kb_dir = Path(str(kb_dir))

    documents: list[Document] = []
    if not kb_dir.exists():
        return documents

    for file_path in sorted(kb_dir.iterdir()):
        if file_path.is_file() and file_path.suffix.lower() in SUPPORTED:
            documents.extend(load_file(file_path))

    return documents


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))